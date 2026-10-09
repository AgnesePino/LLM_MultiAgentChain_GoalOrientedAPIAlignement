"""Generate the project ground-truth registry from the annotation workbook.

The generated module keeps the same public variables used by the notebooks:
SIA_PROJECT_*, ALL_GROUNDTRUTHS, and GROUNDTRUTH_BY_NAME.
"""

from __future__ import annotations

import argparse
import pprint
import re
import textwrap
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL_PATH = Path(__file__).resolve().parent / "GORE_groundtruth_annotation.xlsx"
DEFAULT_OUTPUT_PATH = ROOT / "notebook" / "groundtruth.py"

SKIPPED_SHEETS = {
    "Annotation guide",
}

PROJECT_NAME_BY_SHEET = {
    "SIA_PROJECT_25_26": "SIA Project 25 26",
    "SIA_PROJECT_24_25": "Assegno Unico Universale - SIA Project 24 25",
    "SIA_PROJECT_23_24": "La Reine Marlene - SIA Project 23 24",
    "SIA_PROJECT_22_23": "Ethical Purchasing Group - SIA Project 22 23",
    "SIA_PROJECT_21_22": "Event Organization Portal - SIA Project 21 22",
    "Genome": "Genome Nexus",
    "Gestao Hospital": "Gestao Hospital",
    "London Ambulance Service": "London Ambulance Service",
    "Calculator": "Calculator",
    "Fitnesstracker": "Fitness Tracker",
    "PlanningPoker": "Planning Poker",
    "Recycling": "Recycling",
}

EXTRA_METADATA_BY_NAME = {
    "Genome Nexus": {
        "link-readme": "https://github.com/WebFuzzing/EMB/tree/master/jdk_8_maven/cs/rest-gui/genome-nexus#readme",
        "swagger": "https://raw.githubusercontent.com/WebFuzzing/EMB/refs/heads/master/openapi-swagger/genome-nexus.json",
    },
    "Gestao Hospital": {
        "link-readme": "https://github.com/ValchanOficial/GestaoHospital/blob/master/README.md",
        "swagger": "https://raw.githubusercontent.com/WebFuzzing/EMB/refs/heads/master/openapi-swagger/gestaohospital-rest.json",
    },
}

FIELD_CONFIG = {
    "actors": {"status_col": "B", "value_col": "C"},
    "highLevelGoals": {"status_col": "E", "value_col": "F"},
    "lowLevelGoals": {"status_col": "H", "value_col": "I"},
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = textwrap.dedent(str(value)).strip()
    text = re.sub(r",$", "", text).strip()
    text = text.strip('"').strip()
    return text


def include_status(status: Any, *, only_revision_finished: bool) -> bool:
    normalized = clean_text(status).casefold()
    if normalized == "revision finished":
        return True
    if only_revision_finished:
        return False
    return bool(normalized)


def collect_values(sheet, *, status_col: str, value_col: str, only_revision_finished: bool) -> list[str]:
    values: list[str] = []
    seen: set[str] = set()

    for row in range(7, sheet.max_row + 1):
        if not include_status(sheet[f"{status_col}{row}"].value, only_revision_finished=only_revision_finished):
            continue

        value = clean_text(sheet[f"{value_col}{row}"].value)
        if not value:
            continue

        dedupe_key = value.casefold()
        if dedupe_key not in seen:
            values.append(value)
            seen.add(dedupe_key)

    return values


def variable_name(sheet_name: str) -> str:
    if sheet_name.startswith("SIA_PROJECT_"):
        return sheet_name
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", sheet_name).strip("_").upper()
    return f"SIA_PROJECT_{normalized}"


def build_groundtruths(excel_path: Path, *, only_revision_finished: bool) -> list[tuple[str, dict[str, Any]]]:
    workbook = load_workbook(excel_path, data_only=True)
    projects: list[tuple[str, dict[str, Any]]] = []
    description_sources: dict[str, str] = {}

    for sheet_name in workbook.sheetnames:
        if sheet_name in SKIPPED_SHEETS:
            continue

        sheet = workbook[sheet_name]
        description = clean_text(sheet["B1"].value)
        if not description:
            continue

        description_key = " ".join(description.casefold().split())
        previous_sheet = description_sources.get(description_key)
        if previous_sheet is not None:
            raise ValueError(
                "Duplicate project description in workbook sheets "
                f"'{previous_sheet}' and '{sheet_name}'."
            )
        description_sources[description_key] = sheet_name

        name = PROJECT_NAME_BY_SHEET.get(sheet_name, sheet_name)
        project: dict[str, Any] = {
            "name": name,
            "description": description,
        }
        project.update(EXTRA_METADATA_BY_NAME.get(name, {}))

        for field, config in FIELD_CONFIG.items():
            project[field] = collect_values(
                sheet,
                status_col=config["status_col"],
                value_col=config["value_col"],
                only_revision_finished=only_revision_finished,
            )

        projects.append((variable_name(sheet_name), project))

    return projects


def render_module(projects: list[tuple[str, dict[str, Any]]], *, source_path: Path) -> str:
    lines = [
        '"""Ground-truth registry generated from the GORE annotation workbook.',
        "",
        f"Source workbook: {source_path.as_posix()}",
        "Regenerate with: python notebook/update_groundtruth_from_excel.py",
        '"""',
        "",
        "# This file is generated. Edit the Excel workbook, then regenerate this module.",
        "",
    ]

    printer = pprint.PrettyPrinter(indent=4, width=100, sort_dicts=False)
    for var_name, project in projects:
        lines.append(f"{var_name} = {printer.pformat(project)}")
        lines.append("")

    lines.extend(
        [
            "# Canonical dataset registry shared by the execution and evaluation notebooks.",
            "_REQUIRED_GROUNDTRUTH_FIELDS = {",
            '    "name",',
            '    "description",',
            '    "actors",',
            '    "highLevelGoals",',
            '    "lowLevelGoals",',
            "}",
            "",
            "",
            "def _discover_groundtruths(namespace):",
            "    projects = []",
            "    for variable_name, value in namespace.items():",
            '        if not variable_name.startswith("SIA_PROJECT_"):',
            "            continue",
            "        if not isinstance(value, dict):",
            "            raise TypeError(f\"{variable_name} must be a dictionary\")",
            "",
            "        missing_fields = _REQUIRED_GROUNDTRUTH_FIELDS - value.keys()",
            "        if missing_fields:",
            '            missing = ", ".join(sorted(missing_fields))',
            "            raise ValueError(f\"{variable_name} is missing fields: {missing}\")",
            "        projects.append(value)",
            "",
            "    if not projects:",
            '        raise RuntimeError("No SIA_PROJECT_* ground truths were found")',
            "",
            '    names = [project["name"] for project in projects]',
            "    duplicate_names = sorted({name for name in names if names.count(name) > 1})",
            "    if duplicate_names:",
            "        raise ValueError(",
            '            "Ground-truth project names must be unique: "',
            '            + ", ".join(duplicate_names)',
            "        )",
            "",
            "    return projects",
            "",
            "",
            "ALL_GROUNDTRUTHS = _discover_groundtruths(globals())",
            "",
            "GROUNDTRUTH_BY_NAME = {",
            '    groundtruth["name"]: groundtruth',
            "    for groundtruth in ALL_GROUNDTRUTHS",
            "}",
            "",
        ]
    )

    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Regenerate notebook/groundtruth.py from the Excel annotation workbook."
    )
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL_PATH, help="Path to the .xlsx workbook.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH, help="Python module to generate.")
    parser.add_argument(
        "--only-revision-finished",
        action="store_true",
        help="Exclude rows whose status is not 'Revision finished'. By default every annotated row is included.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    projects = build_groundtruths(args.excel, only_revision_finished=args.only_revision_finished)
    rendered = render_module(projects, source_path=args.excel)
    args.output.write_text(rendered, encoding="utf-8")
    print(f"Generated {args.output} with {len(projects)} projects.")


if __name__ == "__main__":
    main()
